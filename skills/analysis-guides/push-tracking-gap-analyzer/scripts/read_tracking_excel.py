#!/usr/bin/env python3
'Read complete tracking workbook rows into structured JSON.'

import json, sys, os
import openpyxl
from collections import defaultdict

def find_header(ws, required):
    for index, row in enumerate(ws.iter_rows(values_only=True), 1):
        columns = {str(value).strip(): i for i, value in enumerate(row) if value is not None}
        if all(name in columns for name in required):
            return index, columns
        if index >= 20:
            break
    raise ValueError(f"{ws.title}: required columns not found: {required}")


def read_property_sheet(wb, title):
    if title not in wb.sheetnames:
        return {}
    ws = wb[title]
    header_row, columns = find_header(ws, ['属性名', '属性显示名', '属性类型'])
    result = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        def cell(label):
            value = row[columns[label]]
            return str(value).strip() if value is not None else ''
        name = cell('属性名')
        if name:
            result[name] = {'display': cell('属性显示名'), 'type': cell('属性类型')}
    return result


def parse_tracking_excel(excel_path):
    """Parse the tracking Excel and return structured data."""
    wb = openpyxl.load_workbook(excel_path)

    # --- #事件数据 sheet ---
    ws = wb['#事件数据']

    header_row, col_map = find_header(ws, ['事件名', '属性名'])

    current_event = None
    events = {}
    prop_name_set = set()

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        event_name = str(row[col_map['事件名']]).strip() if row[col_map['事件名']] else ''
        prop_name = str(row[col_map['属性名']]).strip() if row[col_map['属性名']] else ''

        if event_name:
            current_event = event_name
            if event_name not in events:
                events[event_name] = {
                    'display': str(row[col_map.get('事件显示名', 0)]) if '事件显示名' in col_map and row[col_map['事件显示名']] else '',
                    'desc': str(row[col_map.get('事件说明', 0)]) if '事件说明' in col_map and row[col_map['事件说明']] else '',
                    'tag': str(row[col_map.get('事件标签', 0)]) if '事件标签' in col_map and row[col_map['事件标签']] else '',
                    'props': {}
                }

        if current_event and prop_name:
            prop_type = str(row[col_map['属性类型']]).strip() if '属性类型' in col_map and row[col_map['属性类型']] else ''
            prop_display = str(row[col_map['属性显示名']]).strip() if '属性显示名' in col_map and row[col_map['属性显示名']] else ''
            prop_desc = str(row[col_map['属性说明']]).strip() if '属性说明' in col_map and row[col_map['属性说明']] else ''
            events[current_event]['props'][prop_name] = {
                'display': prop_display,
                'type': prop_type,
                'desc': prop_desc
            }
            prop_name_set.add(prop_name)

    # Tag distribution
    tags = defaultdict(list)
    for evt, info in events.items():
        tags[info['tag']].append(evt)

    public_props = read_property_sheet(wb, '#公共事件属性')
    user_props = read_property_sheet(wb, '#用户数据')
    prop_name_set.update(public_props)
    prop_name_set.update(user_props)
    wb.close()

    return {
        'events': events,
        'event_count': len(events),
        'all_property_names': sorted(prop_name_set),
        'tag_distribution': {tag: len(evts) for tag, evts in tags.items()},
        'public_props': public_props,
        'user_props': user_props,
    }


if __name__ == '__main__':
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument('--input', required=True, help='Excel file path')
    p.add_argument('--output', required=True, help='Output JSON path')
    args = p.parse_args()

    try:
        result = parse_tracking_excel(args.input)
        os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
        with open(args.output, 'w') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(json.dumps({'status': 'SUCCESS', 'event_count': result['event_count'],
                          'prop_count': len(result['all_property_names'])}))
    except Exception as e:
        print(json.dumps({'status': 'FATAL', 'message': str(e)}))
        sys.exit(1)
