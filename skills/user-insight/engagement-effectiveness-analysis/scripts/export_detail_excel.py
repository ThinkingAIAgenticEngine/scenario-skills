#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Step 4 -- export the detail Excel workbook (shared by both modes)

Reads tmp/analysis.json (produced by Step 3) and tmp/raw_*.json (the raw reports
from Step 2) and builds the 7-sheet detail workbook defined by
detail_workbook_spec.

Design rules (see assets/detail_workbook_spec.md):
  - dump everything; no filtering and no premature aggregation
  - sheets missing in mode B are not deleted; an explanation row names the reason
  - column names match the platform / tracking-plan wording verbatim (the
    original metric keys are preserved)
  - the workbook must be produced by this script, never assembled by hand

Usage:
    python3 export_detail_excel.py --analysis tmp/analysis.json --raw-dir tmp \
        --out-dir output --strategy-name "shiyan"
"""
import argparse
import json
import os
import re
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _wording import _L, _F  # noqa: E402
import _env  # noqa: E402

# openpyxl is the only third-party dependency in this skill, and this is the
# step that needs it. A bare ImportError here is useless to whoever receives the
# skill: it says nothing about which interpreter to use or how to fix it.
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    _env.die_missing_package("openpyxl", "export_detail_excel.py")

# ---------------------------------------------------------------- style constants

HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
HDR_FONT = Font(bold=True, size=10)
TITLE_FONT = Font(bold=True, size=12)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(italic=True, size=10, color="7F6000")
BOLD = Font(bold=True, size=10)
NORMAL = Font(size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)


def log(msg):
    print(f"[excel] {msg}", file=sys.stderr)


def safe_filename(s):
    s = re.sub(r'[\\/:*?"<>|\n\r\t]', "_", str(s or ""))
    return s.strip().strip(".")[:60] or "unnamed"


def write_header(ws, headers, row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 22


def write_row(ws, row, values, bold=False):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = BOLD if bold else NORMAL
        cell.border = BORDER
        cell.alignment = CENTER if not isinstance(v, str) or len(str(v)) < 30 else LEFT
    return row + 1


def write_note(ws, row, text):
    """Write one explanation row (notes why this dimension is unavailable)."""
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = NOTE_FILL
    cell.font = NOTE_FONT
    cell.alignment = LEFT
    return row + 1


def autofit(ws, min_w=10, max_w=46):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            # CJK characters are estimated at 2 columns wide
            txt = str(cell.value)
            w = sum(2 if ord(ch) > 127 else 1 for ch in txt)
            widths[col] = max(widths.get(col, 0), w)
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(max(w + 2, min_w), max_w)


def pct(v, digits=2):
    if v is None:
        return "—"
    try:
        return f"{v * 100:.{digits}f}%"
    except (TypeError, ValueError):
        return "—"


def num(v):
    return "—" if v is None else v


# ---------------------------------------------------------------- sheet builders

def sheet_overview(wb, a):
    """Sheet 1 -- strategy overview"""
    ws = wb.create_sheet(_L("ex.001"))
    m = a["meta"]
    ws["A1"] = _L("ex.002")
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    row = 3
    write_header(ws, [_L("ex.003"), _L("ex.004")], row=row)
    row += 1

    ci = m.get("conversion_indicators") or []
    ci_text = "; ".join(
        f"{c.get('event_name')}（{c.get('event_desc') or '—'}）"
        + (_F("ex.t001", a1=c['window']) if c.get("window") else "")
        + (_F("ex.t002", a1=c['condition']) if c.get("condition") else "")
        for c in ci
    ) or _L("ex.005")

    rows = [
        (_L("ex.006"), m.get("strategy_name")),
        (_L("ex.007"), m.get("task_id")),
        (_L("ex.008"), m.get("project_id")),
        (_L("ex.009"), m.get("mode_label")),
        (_L("ex.010"), m.get("channel")),
        (_L("ex.011"), m.get("status")),
        (_L("ex.012"), m.get("trigger_time_label") or m.get("trigger_time")),
        (_L("ex.013"), m.get("trigger_rule") or "—"),
        (_L("ex.014"), m.get("timezone")),
        (_L("ex.015"), m.get("creator")),
        (_L("ex.016"), m.get("target_cluster") or "—"),
        (_L("ex.017"), _L("ex.018") if m.get("mode") == "A"
                    else _L("ex.019")),
        (_L("ex.020"), ci_text),
        (_L("ex.021"), "; ".join(c["window"] for c in ci if c.get("window")) or "—"),
        (_L("ex.022"), _L("ex.023") if m.get("experiment_enabled") else _L("ex.024")),
        (_L("ex.025"), _L("ex.023") if m.get("frequency_limit_enabled") else _L("ex.024")),
        (_L("ex.026"), _F("ex.t003", a1=a['meta']['data_range'].get('start'), a2=a['meta']['data_range'].get('end'), a3=a['meta']['data_range'].get('days')) if m.get("data_range", {}).get("start") else "—"),
        (_L("ex.027"), "、".join(m.get("blocks_present") or []) or "—"),
        (_L("ex.028"), m.get("generated_at")),
    ]
    for k, v in rows:
        ws.cell(row=row, column=1, value=k).font = BOLD
        ws.cell(row=row, column=1).border = BORDER
        c = ws.cell(row=row, column=2, value=v if v is not None else "—")
        c.border = BORDER
        c.font = NORMAL
        c.alignment = LEFT
        row += 1

    row += 1
    ws.cell(row=row, column=1, value=_L("ex.029")).font = TITLE_FONT
    row += 1
    write_header(ws, [_L("ex.030"), _L("ex.031"), _L("ex.032")], row=row)
    row += 1
    if a["degradations"]:
        for d in a["degradations"]:
            row = write_row(ws, row, [d["item"], d["impact"], d["loss"]])
    else:
        row = write_row(ws, row, [_L("ex.033"), _L("ex.034"), "0%"])
    row += 1
    ws.cell(row=row, column=1,
            value=_L("ex.035")).font = NOTE_FONT

    # Caliber note: one business volume is measured differently per endpoint, so
    # stating it explicitly keeps readers from miscalculating.
    cavs = a.get("field_caveats") or []
    if cavs:
        row += 2
        ws.cell(row=row, column=1, value=_L("ex.036")).font = TITLE_FONT
        row += 1
        write_header(ws, [_L("ex.037"), _L("ex.038")], row=row)
        row += 1
        for c in cavs:
            row = write_row(ws, row, [
                c.get("field"),
                f"{c.get('source')} = {c.get('value')}；{c.get('counterpart')}。{c.get('note')}",
            ])
        row += 1
        ws.cell(row=row, column=1,
                value=_L("ex.039")).font = NOTE_FONT

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 78
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A3"
    return ws


def sheet_funnel(wb, a):
    """Sheet 2 -- funnel"""
    ws = wb.create_sheet(_L("ex.040"))
    ws["A1"] = _L("ex.041")
    ws["A1"].font = TITLE_FONT
    row = 3
    write_header(ws, [_L("ex.042"), _L("ex.043"), _L("ex.044"), _L("ex.045"), _L("ex.046"),
                      _L("ex.047"), _L("ex.048"), _L("ex.049"), _L("ex.050")], row=row)
    row += 1

    brk = a.get("data_break") or {}
    brk_key = brk.get("metric_key")
    for s in a["funnel"]:
        note = ""
        if s["metric_key"] == brk_key:
            note = _L("ex.051")
        elif s["metric_key"] in (brk.get("downstream_stages") or []):
            note = _L("ex.052")
        row = write_row(ws, row, [
            s["stage"], s["metric_key"], s["desc"], s["value"],
            pct(s["stage_conv"]), pct(s["overall_conv"]),
            num(s["drop"]), pct(s["drop_rate"]), note,
        ])

    row += 1
    ws.cell(row=row, column=1, value=_L("ex.053")).font = BOLD
    mb = a.get("max_bleed")
    ws.cell(row=row, column=2,
            value=mb["reason"] if mb else _L("ex.054")).font = NORMAL
    row += 1
    ws.cell(row=row, column=1, value=_L("ex.055")).font = BOLD
    ws.cell(row=row, column=2,
            value=_L("ex.056")).font = NORMAL
    row += 2

    ws.cell(row=row, column=1, value=_L("ex.057")).font = TITLE_FONT
    row += 1
    write_header(ws, [_L("ex.058"), _L("ex.059"), _L("ex.060"), _L("ex.061"), _L("ex.062"), _L("ex.050")], row=row)
    row += 1
    for r in a["rates"]:
        note = _F("ex.t004", a1=r['denominator_metric']) if r.get("degraded") else ""
        row = write_row(ws, row, [
            r["name"], pct(r["value"]), pct(r["threshold"]) if r["threshold"] else "—",
            f"{r['flag']} {r['advice']}", r["desc"], note,
        ])
    autofit(ws)
    ws.freeze_panes = "A4"
    return ws


def sheet_dimension(wb, a):
    """Sheet 3 -- dimension detail (per send instance x metric)"""
    ws = wb.create_sheet(_L("ex.063"))
    ws["A1"] = _L("ex.064")
    ws["A1"].font = TITLE_FONT
    row = 3

    inst = a.get("instances") or []
    if not inst:
        row = write_note(ws, row, _L("ex.065"))
        autofit(ws)
        return ws

    keys = [k for k in (a["metrics"] or {}).keys()]
    headers = [_L("ex.066")] + [f"{k}" for k in keys]
    write_header(ws, headers, row=row)
    row += 1
    for it in inst:
        row = write_row(ws, row, [it.get("instance_id")] + [it.get(k, 0) for k in keys])

    row += 1
    row = write_note(ws, row,
                     _L("ex.067"))
    autofit(ws)
    ws.freeze_panes = "B4"
    return ws


def sheet_timeline(wb, a):
    """Sheet 4 -- time series"""
    ws = wb.create_sheet(_L("ex.068"))
    ws["A1"] = _L("ex.069")
    ws["A1"].font = TITLE_FONT
    row = 3

    tl = a.get("timeline") or []
    if not tl:
        row = write_note(ws, row, _L("ex.070"))
        autofit(ws)
        return ws

    keys = [k for k in (a["metrics"] or {}).keys()]
    write_header(ws, [_L("ex.071")] + keys, row=row)
    row += 1
    first_data_row = row
    for r in tl:
        row = write_row(ws, row, [r.get("date")] + [r.get(k, 0) for k in keys])
    last_data_row = row - 1

    # summary row
    tr = a.get("timeline_total") or a.get("metrics") or {}
    row = write_row(ws, row, [_L("ex.072")] + [tr.get(k, 0) for k in keys], bold=True)
    for c in range(1, len(keys) + 2):
        ws.cell(row=row - 1, column=c).fill = PatternFill("solid", fgColor="EDEDED")

    row += 1
    row = write_note(ws, row,
                     _F("ex.t005", a1=len(tl)))
    autofit(ws)
    ws.freeze_panes = "B4"
    return ws


def sheet_segments(wb, a):
    """Sheet 5 -- audience breakdown"""
    ws = wb.create_sheet(_L("ex.073"))
    ws["A1"] = _L("ex.074")
    ws["A1"].font = TITLE_FONT
    row = 3

    seg = a.get("segments") or []
    if not seg:
        row = write_note(ws, row, _L("ex.075"))
        row += 1
        row = write_note(ws, row,
                         _L("ex.076"))
        autofit(ws)
        return ws

    write_header(ws, [_L("ex.077"), _L("ex.078"), _L("ex.079"), _L("ex.080"), _L("ex.081"), _L("ex.082")], row=row)
    row += 1
    for s in seg:
        row = write_row(ws, row, [s.get("dimension"), s.get("value"), s.get("reach"),
                                  s.get("click"), s.get("convert"), pct(s.get("response_rate"))])
    autofit(ws)
    return ws


def sheet_experiment(wb, a):
    """Sheet 6 -- strategy and group comparison"""
    ws = wb.create_sheet(_L("ex.083"))
    ws["A1"] = _L("ex.084")
    ws["A1"].font = TITLE_FONT
    row = 3

    exp = a.get("experiment")
    if not exp or not exp.get("groups"):
        row = write_note(ws, row, _L("ex.085"))
        autofit(ws)
        return ws

    write_header(ws, [_L("ex.086"), _L("ex.087"), _L("ex.088"), _L("ex.089"),
                      _L("ex.090"), _L("ex.091"), _L("ex.092"), _L("ex.093"), _L("ex.094")], row=row)
    row += 1
    for g in exp["groups"]:
        row = write_row(ws, row, [
            g.get("group_id"), g.get("group_name"), g.get("group_role"), g.get("success_user_num"),
            g.get("lift_rate"), g.get("ci_floor"), g.get("ci_ceil"),
            g.get("significance_level"), {True: _L("ex.023"), False: _L("ex.024"), None: "—"}.get(g.get("win"), "—"),
        ])

    row += 1
    ws.cell(row=row, column=1, value=_L("ex.095")).font = BOLD
    ws.cell(row=row, column=2, value=f"{exp.get('verdict') or '—'}：{exp.get('conclusion') or '—'}").font = NORMAL
    row += 2

    mx = exp.get("matrix")
    if mx:
        ws.cell(row=row, column=1, value=_L("ex.096")).font = TITLE_FONT
        row += 1
        keys = [k for k in (a["metrics"] or {}).keys()]
        write_header(ws, [_L("ex.071"), _L("ex.066"), _L("ex.097"), _L("ex.086")] + keys, row=row)
        row += 1
        for r in mx:
            row = write_row(ws, row, [r.get("date"), r.get("instance_id"), r.get("group_name"),
                                      r.get("group_id")] + [r.get(k) for k in keys])
    autofit(ws)
    return ws


def sheet_raw(wb, a, raw_dir):
    """Sheet 7 -- raw data (endpoint response fields flattened)"""
    ws = wb.create_sheet(_L("ex.098"))
    ws["A1"] = _L("ex.099")
    ws["A1"].font = TITLE_FONT
    row = 3

    write_header(ws, [_L("ex.100"), _L("ex.037"), _L("ex.101")], row=row)
    row += 1

    # raw overview metric values
    for k, v in (a.get("metrics") or {}).items():
        row = write_row(ws, row, ["task_data_overview", k, v])

    # task metadata
    p = os.path.join(raw_dir, "raw_tasklist.json")
    if os.path.exists(p):
        try:
            raw = json.load(open(p, encoding="utf-8"))
            ts = ((raw.get("data") or {}).get("operationTaskDTOS")) or []
            tid = a["meta"].get("task_id")
            t = next((x for x in ts if str(x.get("taskId")) == str(tid)), None)
            if t:
                for k, v in t.items():
                    sv = json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v
                    row = write_row(ws, row, ["task_list", k, sv])
        except Exception as e:
            row = write_note(ws, row, _F("ex.t006", a1=e))

    # raw experiment payload
    p = os.path.join(raw_dir, "raw_exp_overview.json")
    if os.path.exists(p):
        try:
            raw = json.load(open(p, encoding="utf-8"))
            for g in ((raw.get("data") or {}).get("overviewData") or []):
                for k, v in g.items():
                    row = write_row(ws, row, ["task_experiment_report", f"{g.get('expGroupName')}.{k}",
                                              json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
        except Exception as e:
            row = write_note(ws, row, _F("ex.t007", a1=e))

    row += 1
    row = write_note(ws, row, _L("ex.102"))
    for c in (a.get("field_caveats") or []):
        row = write_note(ws, row,
                         _F("ex.t008", a1=c.get('field'), a2=c.get('value'), a3=c.get('source'), a4=c.get('counterpart'), a5=c.get('note')))
    autofit(ws)
    ws.column_dimensions["C"].width = 70
    ws.freeze_panes = "A4"
    return ws


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=_L("ex.103"))
    ap.add_argument("--analysis", default="tmp/analysis.json")
    ap.add_argument("--raw-dir", default="tmp")
    ap.add_argument("--out-dir", default="output")
    ap.add_argument("--strategy-name", default=None)
    args = ap.parse_args()

    with open(args.analysis, encoding="utf-8") as f:
        a = json.load(f)

    strategy = args.strategy_name or a["meta"].get("strategy_name") or a["meta"].get("task_id") or "strategy"
    date = datetime.now().strftime("%Y%m%d")
    os.makedirs(args.out_dir, exist_ok=True)
    out_path = os.path.join(args.out_dir, _F("ex.t009", a1=safe_filename(strategy), a2=date))

    wb = Workbook()
    wb.remove(wb.active)          # 删掉默认空表

    sheet_overview(wb, a)
    sheet_funnel(wb, a)
    sheet_dimension(wb, a)
    sheet_timeline(wb, a)
    sheet_segments(wb, a)
    sheet_experiment(wb, a)
    sheet_raw(wb, a, args.raw_dir)

    wb.save(out_path)
    log(_F("ex.t010", a1=out_path))
    log(f"   sheet: {wb.sheetnames}")
    print(out_path)


if __name__ == "__main__":
    main()
