#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Step 3-5 deliverable gate

Checks all three deliverables against the delivery standard in a single pass:
  - analysis.json  : every required analysis block present
  - detail Excel   : all 7 worksheets present + each sheet has more than 0 rows
  - analysis report: file non-empty + all six sections present + data-source
                     mode stated + confidence statement present

Usage:
    python3 validate_outputs.py --analysis tmp/analysis.json \
        --excel output/xxx_detail_20260914.xlsx \
        --report output/xxx_effectiveness_report_20260914.md

Exit code: 0 = everything passed; 1 = at least one failure
"""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _wording import _L, _F  # noqa: E402
import _env  # noqa: E402

try:
    from openpyxl import load_workbook
except ImportError:
    _env.die_missing_package("openpyxl", "validate_outputs.py")

SECTIONS = [_L("vo.001"), _L("vo.002"), _L("vo.003"), _L("vo.004"), _L("vo.005"), _L("vo.006")]
REQUIRED_SHEETS = [
    _L("vo.007"), _L("vo.008"), _L("vo.009"), _L("vo.010"),
    _L("vo.011"), _L("vo.012"), _L("vo.013"),
]
REQUIRED_ANALYSIS_BLOCKS = ["meta", "metrics", "funnel", "rates", "issues",
                            "recommendations", "validation_plan", "degradations"]

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--analysis", default="tmp/analysis.json")
    ap.add_argument("--excel", required=True)
    ap.add_argument("--report", required=True)
    args = ap.parse_args()

    # ---------------- analysis.json
    if not os.path.exists(args.analysis):
        check(_L("vo.014"), False, args.analysis)
    else:
        check(_L("vo.014"), True)
        a = json.load(open(args.analysis, encoding="utf-8"))
        missing = [k for k in REQUIRED_ANALYSIS_BLOCKS if k not in a]
        check(_L("vo.015"), not missing, _F("vo.t001", a1=missing) if missing else _F("vo.t002", a1=len(REQUIRED_ANALYSIS_BLOCKS)))
        check(_L("vo.016"), len(a.get("funnel") or []) >= 2, _F("vo.t003", a1=len(a.get('funnel') or [])))
        check(_L("vo.017"), len(a.get("rates") or []) >= 1, _F("vo.t004", a1=len(a.get('rates') or [])))
        check(_L("vo.018"), len(a.get("issues") or []) >= 1, _F("vo.t005", a1=len(a.get('issues') or [])))
        check(_L("vo.019"), bool(a.get("meta", {}).get("mode_label")),
              str(a.get("meta", {}).get("mode_label")))
        check(_L("vo.020"), isinstance(a.get("degradations"), list),
              _F("vo.t006", a1=len(a.get('degradations') or [])))

    # ---------------- Excel
    if not os.path.exists(args.excel):
        check(_L("vo.021"), False, args.excel)
    else:
        check(_L("vo.021"), True, f"{os.path.getsize(args.excel)} bytes")
        wb = load_workbook(args.excel)
        missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
        check(_L("vo.022"), not missing, _F("vo.t001", a1=missing) if missing else _L("vo.023"))
        empty = []
        for s in REQUIRED_SHEETS:
            if s not in wb.sheetnames:
                continue
            ws = wb[s]
            n = sum(1 for row in ws.iter_rows() for c in row if c.value is not None)
            if n == 0:
                empty.append(s)
        check(_L("vo.024"), not empty, _F("vo.t007", a1=empty) if empty else _L("vo.025"))

    # ---------------- report
    if not os.path.exists(args.report):
        check(_L("vo.026"), False, args.report)
    else:
        size = os.path.getsize(args.report)
        md = open(args.report, encoding="utf-8").read()
        check(_L("vo.026"), True, f"{size} bytes")
        check(_L("vo.027"), size > 500, f"{size} bytes")
        miss_sec = [s for s in SECTIONS if s not in md]
        check(_L("vo.028"), not miss_sec, _F("vo.t001", a1=miss_sec) if miss_sec else _L("vo.029"))
        check(_L("vo.030"), _L("vo.031") in md)
        check(_L("vo.032"), _L("vo.033") in md)
        check(_L("vo.034"), "Sheet" in md)

    # ---------------- output
    print("=" * 62)
    print(_L("vo.035"))
    print("=" * 62)
    failed = 0
    for name, ok, detail in results:
        mark = "✅" if ok else "❌"
        if not ok:
            failed += 1
        print(f"  {mark} {name}" + (f"  —— {detail}" if detail else ""))
    print("=" * 62)
    print(_F("vo.t008", a1=len(results), a2=len(results) - failed, a3=failed))
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
